#!/usr/bin/env python3
"""Write docs/SpriteChecklist.xlsx: which sprites we have for all 1025 species (and every Mega).

A check mark means we have the sprite; a highlighted empty cell means we still need it.
  * #1-493 already exist in Renegade Platinum (the owner's rule: keep Renegade's sprites).
  * #494-1025 backs come from data/sprite_index.csv (the DPPt Style Gen VI and Beyond pack).
  * Fronts and box/party icons for added species: none yet.
Re-run after adding sprites (and after re-running tools/sprites/import_sprite_pack.py).

Usage: python3 tools/sprites/gen_sprite_checklist.py
"""
import csv
import os

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
CHECK = "✓"
FONT = "Arial"
MISSING = PatternFill("solid", start_color="F4CCCC")
HEADER = PatternFill("solid", start_color="1F3864")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
SPRITE_COLS = ["Non-shiny front sprite", "Non-shiny back sprite", "Shiny front sprite", "Shiny back sprite",
               "Box/party screen sprite"]
# Owner-supplied Mega fronts (assets/sprites/owner_supplied/).
OWNER_MEGA_FRONTS = {"Mega Garchomp Z": "Owner-supplied front + shiny front; 116×95 art must be redrawn to fit the DS 80×80 frame."}


def load():
    species = list(csv.DictReader(open(os.path.join(ROOT, "data", "species.csv"), encoding="utf-8")))
    megas = list(csv.DictReader(open(os.path.join(ROOT, "data", "megas.csv"), encoding="utf-8")))
    index = list(csv.DictReader(open(os.path.join(ROOT, "data", "sprite_index.csv"), encoding="utf-8")))
    files = {}
    for e in index:
        if e["category"] in ("base", "mega") and "(female)" not in e["form"]:
            files[(e["form"], e["kind"])] = e
    return species, megas, files


def pack_note(entries):
    notes = []
    for label, e in entries:
        if e and e["ds_ready"] != "1":
            why = f"{e['colours']} colours" if int(e["colours"]) > 15 else "not a clean 2× upscale"
            notes.append(f"{label}: {why}")
    return ("Needs DS conversion — " + "; ".join(notes)) if notes else ""


def style_header(ws, headers):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "C2"


def write_row(ws, r, values, sprite_start, n_sprites):
    for c, v in enumerate(values, 1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = Font(name=FONT)
        cell.border = BORDER
        if sprite_start <= c < sprite_start + n_sprites:
            cell.alignment = Alignment(horizontal="center")
            if v != CHECK:
                cell.fill = MISSING


def main():
    species, megas, files = load()
    wb = Workbook()

    # ---- Sheet 1: all 1025 species
    ws = wb.active
    ws.title = "Sprite Checklist"
    headers = ["Dex #", "Pokémon", "Generation", *SPRITE_COLS, "Source", "Notes"]
    style_header(ws, headers)
    for i, s in enumerate(species, start=2):
        if s["in_renegade"] == "1":
            marks = [CHECK] * 5
            source = "Renegade Platinum ROM (existing sprites)"
            note = ""
        else:
            back = files.get((s["name"], "back"))
            sback = files.get((s["name"], "back_shiny"))
            marks = ["", CHECK if back else "", "", CHECK if sback else "", ""]
            source = "DPPt Style Gen VI and Beyond pack (backs)" if (back or sback) else ""
            note = pack_note([("back", back), ("shiny back", sback)])
        write_row(ws, i, [int(s["dex"]), s["name"], int(s["generation"]), *marks, source, note], 4, 5)
    last = len(species) + 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"
    for col, w in zip("ABCDEFGHIJ", [7, 18, 11, 14, 14, 14, 14, 14, 38, 50]):
        ws.column_dimensions[col].width = w

    # ---- Sheet 2: Mega Evolutions / Primal Reversions
    wm = wb.create_sheet("Mega & Primal Forms")
    mheaders = ["Dex #", "Form", "Non-shiny front sprite", "Non-shiny back sprite", "Shiny front sprite",
                "Shiny back sprite", "Source", "Notes"]
    style_header(wm, mheaders)
    wm.cell(row=1, column=2).comment = Comment(
        "Mega/Primal forms have no box/party column: the party and PC show the base form's icon.", "Imperial Platinum")
    for i, m in enumerate(megas, start=2):
        back = files.get((m["form"], "back"))
        sback = files.get((m["form"], "back_shiny"))
        front = m["form"] in OWNER_MEGA_FRONTS
        marks = [CHECK if front else "", CHECK if back else "", CHECK if front else "", CHECK if sback else ""]
        srcs = []
        if back or sback:
            srcs.append("DPPt Style Gen VI and Beyond pack (backs)")
        if front:
            srcs.append("Owner-supplied (fronts)")
        note = "; ".join(x for x in (OWNER_MEGA_FRONTS.get(m["form"], ""), pack_note([("back", back), ("shiny back", sback)])) if x)
        write_row(wm, i, [int(m["dex"]), m["form"], *marks, "; ".join(srcs), note], 3, 4)
    mlast = len(megas) + 1
    wm.auto_filter.ref = f"A1:H{mlast}"
    for col, w in zip("ABCDEFGH", [7, 32, 14, 14, 14, 14, 44, 60]):
        wm.column_dimensions[col].width = w

    # ---- Sheet 3: summary (live formulas)
    su = wb.create_sheet("Summary", 0)
    su["A1"] = "Imperial Platinum — Sprite Aggregation Summary"
    su["A1"].font = Font(name=FONT, bold=True, size=14)
    su["A2"] = "✓ = we have the sprite. Red-highlighted empty cell = still needed. Counts update as you add ✓ marks."
    su["A3"] = ("#1–493 are already in Renegade Platinum (we keep its sprites). Added-species backs come from the "
                "DPPt Style Gen VI and Beyond pack; no added-species fronts or box icons yet.")
    for r in (2, 3):
        su[f"A{r}"].font = Font(name=FONT, italic=True)
    hdr = ["Sprite", "All species: have", "All species: missing", "Added species (#494–1025): have",
           "Added species: missing", "Megas: have", "Megas: missing"]
    for c, h in enumerate(hdr, 1):
        cell = su.cell(row=5, column=c, value=h)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = BORDER
    sc = "'Sprite Checklist'"
    mg = "'Mega & Primal Forms'"
    added_first = next(i for i, s in enumerate(species, start=2) if s["in_renegade"] == "0")
    mega_col = {SPRITE_COLS[0]: "C", SPRITE_COLS[1]: "D", SPRITE_COLS[2]: "E", SPRITE_COLS[3]: "F"}
    for r, name in enumerate(SPRITE_COLS, start=6):
        col = get_column_letter(4 + SPRITE_COLS.index(name))
        all_rng = f"{sc}!${col}$2:${col}${last}"
        add_rng = f"{sc}!${col}${added_first}:${col}${last}"
        row = [name,
               f'=COUNTIF({all_rng},"{CHECK}")', f'=ROWS({all_rng})-B{r}',
               f'=COUNTIF({add_rng},"{CHECK}")', f'=ROWS({add_rng})-D{r}']
        if name in mega_col:
            mrng = f"{mg}!${mega_col[name]}$2:${mega_col[name]}${mlast}"
            row += [f'=COUNTIF({mrng},"{CHECK}")', f'=ROWS({mrng})-F{r}']
        else:
            row += ["n/a", "n/a"]
        for c, v in enumerate(row, 1):
            cell = su.cell(row=r, column=c, value=v)
            cell.font = Font(name=FONT, bold=(c == 1))
            cell.border = BORDER
            if c > 1:
                cell.alignment = Alignment(horizontal="center")
    su.column_dimensions["A"].width = 26
    for col in "BCDEFG":
        su.column_dimensions[col].width = 17
    su.row_dimensions[5].height = 32

    out = os.path.join(ROOT, "docs", "SpriteChecklist.xlsx")
    wb.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
